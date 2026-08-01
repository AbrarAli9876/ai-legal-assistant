from fastapi import (
    APIRouter,
    HTTPException,
    UploadFile,
    File
)

import pypdf
import docx
import io


# ============================================================
# FIR Analyzer Service
# ============================================================

from app.services.fir_evidence_analyzer import analyze_fir


# ============================================================
# Router
# ============================================================

router = APIRouter(
    prefix="/api/v1/analyzer",
    tags=["FIRAnalyzer"]
)


# ============================================================
# Configuration
# ============================================================

MAX_FILE_SIZE = 20 * 1024 * 1024

MAX_DOCUMENT_CHARS = 400000


# ============================================================
# PDF Text Extraction
# ============================================================

def extract_text_from_pdf(
    file_stream: io.BytesIO
) -> str:

    try:

        reader = pypdf.PdfReader(
            file_stream
        )

        text_parts = []

        for page in reader.pages:

            page_text = (
                page.extract_text()
                or ""
            )

            if page_text:

                text_parts.append(
                    page_text
                )

        text = "\n".join(
            text_parts
        )

        return text.replace(
            "\x00",
            ""
        )

    except Exception as e:

        print(
            "[fir_analyzer] PDF read error:",
            str(e)
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read PDF. "
                "The file may be corrupt, encrypted, "
                f"or image-only: {str(e)}"
            )
        )


# ============================================================
# DOCX Text Extraction
# ============================================================

def extract_text_from_docx(
    file_stream: io.BytesIO
) -> str:

    try:

        document = docx.Document(
            file_stream
        )

        text_parts = []

        # ----------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            paragraph_text = (
                paragraph.text.strip()
            )

            if paragraph_text:

                text_parts.append(
                    paragraph_text
                )

        # ----------------------------------------------------
        # Tables
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_values = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:

                        row_values.append(
                            cell_text
                        )

                if row_values:

                    text_parts.append(
                        " | ".join(
                            row_values
                        )
                    )

        text = "\n".join(
            text_parts
        )

        return text.replace(
            "\x00",
            ""
        )

    except Exception as e:

        print(
            "[fir_analyzer] DOCX read error:",
            str(e)
        )

        raise HTTPException(
            status_code=400,
            detail=(
                f"Could not read DOCX file: {str(e)}"
            )
        )


# ============================================================
# FIR Analysis Endpoint
# ============================================================

@router.post(
    "/analyze-fir"
)
async def handle_fir_analysis(
    file: UploadFile = File(...)
):

    print(
        "\n[fir_analyzer] request start"
    )

    # ========================================================
    # Validate Filename
    # ========================================================

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="No file was provided."
        )

    # ========================================================
    # Read File
    # ========================================================

    try:

        contents = await file.read()

    except Exception as e:

        print(
            "[fir_analyzer] upload read error:",
            str(e)
        )

        raise HTTPException(
            status_code=400,
            detail="Could not read uploaded file."
        )

    # ========================================================
    # Validate Contents
    # ========================================================

    if not contents:

        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    if len(contents) > MAX_FILE_SIZE:

        raise HTTPException(
            status_code=413,
            detail=(
                "File is too large. "
                "Maximum supported size is 20 MB."
            )
        )

    # ========================================================
    # File Information
    # ========================================================

    content_type = (
        file.content_type
        or ""
    ).lower()

    filename = (
        file.filename
        or ""
    ).lower()

    file_stream = io.BytesIO(
        contents
    )

    # ========================================================
    # Extract Text
    # ========================================================

    try:

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        if (
            content_type == "application/pdf"
            or filename.endswith(".pdf")
        ):

            print(
                "[fir_analyzer] reading PDF"
            )

            document_text = (
                extract_text_from_pdf(
                    file_stream
                )
            )

        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        elif (
            "wordprocessingml"
            in content_type
            or filename.endswith(".docx")
        ):

            print(
                "[fir_analyzer] reading DOCX"
            )

            document_text = (
                extract_text_from_docx(
                    file_stream
                )
            )

        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        elif (
            content_type.startswith("text/")
            or filename.endswith(".txt")
        ):

            print(
                "[fir_analyzer] reading TXT"
            )

            try:

                document_text = (
                    contents.decode(
                        "utf-8"
                    )
                )

            except UnicodeDecodeError:

                document_text = (
                    contents.decode(
                        "utf-8",
                        errors="ignore"
                    )
                )

            document_text = (
                document_text.replace(
                    "\x00",
                    ""
                )
            )

        # ----------------------------------------------------
        # Unsupported File
        # ----------------------------------------------------

        else:

            raise HTTPException(
                status_code=400,
                detail=(
                    "Invalid file type. "
                    "Please upload PDF, DOCX, or TXT."
                )
            )

    except HTTPException:
        raise

    except Exception as e:

        print(
            "[fir_analyzer] extraction error:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=(
                f"Error reading file: {str(e)}"
            )
        )

    # ========================================================
    # Validate Extracted Text
    # ========================================================

    document_text = (
        document_text.strip()
    )

    if not document_text:

        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text could be extracted "
                "from the uploaded file. "
                "If this is a scanned PDF, OCR may be required."
            )
        )

    print(
        "[fir_analyzer] extracted characters:",
        len(document_text)
    )

    # ========================================================
    # Limit Document Size
    # ========================================================

    if len(document_text) > MAX_DOCUMENT_CHARS:

        print(
            "[fir_analyzer] truncating document:",
            len(document_text),
            "->",
            MAX_DOCUMENT_CHARS
        )

        document_text = (
            document_text[
                :MAX_DOCUMENT_CHARS
            ]
        )

    # ========================================================
    # Call FIR Analyzer Service
    # ========================================================

    try:

        print(
            "[fir_analyzer] calling analyzer service"
        )

        analysis = await analyze_fir(
            document_text
        )

        print(
            "[fir_analyzer] analysis successful"
        )

        return analysis

    # ========================================================
    # AI Service Errors
    # ========================================================

    except ValueError as e:

        print(
            "[fir_analyzer] AI error:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

    # ========================================================
    # Unexpected Errors
    # ========================================================

    except Exception as e:

        print(
            "[fir_analyzer] unexpected error:",
            type(e).__name__,
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "An unexpected error occurred "
                "while analyzing the FIR."
            )
        )